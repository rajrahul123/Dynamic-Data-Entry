"""Generic record export service: CSV, Excel (.xlsx), PDF, and SQL.

Everything here is generic and data-driven. Exports are rendered from the
live ``FormField`` definitions (columns follow ``sort_order``, headers use the
field label) plus the ``Submission`` rows themselves — there is no form-
specific knowledge anywhere.

The Phase 5 query engine is reused verbatim: ``build_record_query`` produces
the same WHERE clauses and ORDER BY expressions used by the records list, so a
CSV/Excel/PDF/SQL export honors the exact same ``search`` / ``filters`` /
``sort_by`` / ``sort_order`` semantics as the UI, with no second filtering
implementation. Export never applies ``limit`` / ``offset`` from the Records
page; instead it exports every matching record up to
:data:`MAX_EXPORT_RECORDS`.

Export formats:

* ``csv``  — UTF-8 (with BOM so Excel renders Unicode correctly),
  RFC-4180-style escaping via the standard ``csv`` module.
* ``xlsx`` — single ``Records`` worksheet via ``openpyxl``; native numbers
  and datetime values where practical, header formatting, readable widths.
* ``pdf``  — landscape, wrapped cells, repeated header rows, via reportlab.
* ``sql``  — ``INSERT INTO submissions (...) VALUES (...);`` statements for
  PostgreSQL, safely single-quote-escaped; intended for inspection or data
  migration, **not** a complete database backup.
* ``pdf-form`` — portrait, form-style PDF with **one submission per page**
  (Phase 8). Renders field labels and values from the live form definition,
  reusing the same ``_display_value`` formatting as every other format.
  Non-Latin script runs (Devanagari / CJK) are wrapped in registered Unicode
  fonts when any are available on the host, so Hindi and CJK text renders when
  the environment provides suitable fonts. There is **no executable template
  system**: the layout is generated strictly from the form definition.

The SQL export is the *only* format that emits raw ``Submission.data`` JSON;
spreadsheet/PDF exports render only the columns defined by the current form
definition (unknown stored keys are never added as arbitrary columns). The
table PDF additionally caps any single cell at ``_PDF_RECORD_CELL_LIMIT`` so a
pathological value can never abort reportlab; CSV / XLSX / SQL always carry the
complete value and the form-style PDFs apply the same cell cap only to field
values.
"""

from __future__ import annotations

import csv
import io
import json
import os
import re
import string
from datetime import date, datetime, time
from typing import Any, BinaryIO
from xml.sax.saxutils import escape as xml_escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy import and_, func, select
from sqlalchemy.orm import Session, selectinload

from app.models import FieldType, Form, Submission
from app.services.record_query import RecordQueryError, build_record_query

EXPORT_FORMATS = ("csv", "xlsx", "pdf", "sql", "pdf-form")

#: Hard server-side cap on exported rows. Exports are synchronous and bounded;
#: a single export larger than this should be narrowed via search/filters.
#: Exceeding the cap is an error (400) -- never a silent truncation.
MAX_EXPORT_RECORDS = 10000

#: Column headers exported before the dynamic form fields (documented).
EXPORT_META_COLUMNS = ("ID", "Submitted By", "Submitted At")

_MEDIA_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "pdf-form": "application/pdf",
    "sql": "application/sql",
}

#: Characters stripped from a generated filename; the slug is built only from
#: ASCII letters/digits/dashes so it is always a safe attachment filename.
_FILENAME_PATTERN = re.compile(r"[^a-z0-9]+")


class ExportError(ValueError):
    """Raised for invalid export requests (bad format, oversized dataset)."""


class ExportSizeError(ExportError):
    """Raised when the matching records exceed :data:`MAX_EXPORT_RECORDS`."""


def _option_label(field, value: Any) -> str | None:
    """Resolve a select/radio value (stored value *or* label) to its label."""
    if value is None:
        return None
    raw = str(value)
    for option in (field.settings or {}).get("options") or []:
        if not isinstance(option, dict):
            continue
        if str(option.get("value")) == raw or str(option.get("label")) == raw:
            return str(option.get("label") or option.get("value") or raw)
    return None


def _display_value(field, value: Any) -> str:
    """Human-readable cell value shared by CSV/PDF (Excel builds on this)."""
    if value is None:
        return ""
    if field.field_type in (FieldType.select, FieldType.radio):
        return _option_label(field, value) or str(value)
    if field.field_type is FieldType.checkbox:
        return "Yes" if value is True else "No"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return str(value)


def _parse_temporal(field, value: Any) -> date | time | datetime | None:
    """Best-effort parse of stored date/time/datetime text for Excel cells."""
    if not isinstance(value, str):
        return None
    text = value.strip()
    try:
        if field.field_type is FieldType.date:
            return date.fromisoformat(text)
        if field.field_type is FieldType.time:
            return time.fromisoformat(text)
        if field.field_type is FieldType.datetime:
            parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
            return parsed.replace(tzinfo=None)
    except ValueError:
        return None
    return None


def _excel_value(field, value: Any) -> Any:
    """Native value for an Excel cell (numbers/datetimes, else display text)."""
    if value is None:
        return None
    if field.field_type is FieldType.number and isinstance(value, (int, float)) and not isinstance(
        value, bool
    ):
        return value
    if field.field_type in (FieldType.date, FieldType.time, FieldType.datetime):
        parsed = _parse_temporal(field, value)
        if parsed is not None:
            return parsed
    return _display_value(field, value)


def _rows(form: Form, submissions: list[Submission]) -> tuple[list[str], list[list[str]]]:
    """Build the header row and display rows for spreadsheet/PDF exports."""
    fields = sorted(form.fields, key=lambda f: f.sort_order)

    def row_for(submission: Submission) -> list[str]:
        submitter = submission.submitted_by_user
        submitter_name = (
            submitter.full_name
            if submitter and submitter.full_name
            else (submitter.username if submitter else "")
        )
        cells: list[str] = [str(submission.id), submitter_name or "", submission.submitted_at.isoformat() if submission.submitted_at else ""]
        for field in fields:
            cells.append(_display_value(field, submission.data.get(field.field_key)))
        return cells

    headers = [*EXPORT_META_COLUMNS, *[field.label for field in fields]]
    return headers, [row_for(submission) for submission in submissions]


def slugify_filename(form_name: str, form_id: int) -> str:
    """Build a safe, useful filename stem from the form name."""
    slug = _FILENAME_PATTERN.sub("-", form_name.lower()).strip("-").strip(string.punctuation)
    return slug or f"form-{form_id}"


def _fetch_records(
    db: Session,
    form: Form,
    search: str | None,
    filters: str | None,
    sort_by: str | None,
    sort_order: str | None,
) -> list[Submission]:
    """Query matching submissions with the Phase 5 engine (limit enforced)."""
    where, order_by = build_record_query(
        db,
        form,
        search=search,
        filters=filters,
        sort_by=sort_by,
        sort_order=sort_order,
    )
    total_where = and_(*where)
    total = (
        db.scalar(select(func.count()).select_from(Submission).where(total_where)) or 0
    )
    if total > MAX_EXPORT_RECORDS:
        raise ExportSizeError(
            f"Export is limited to {MAX_EXPORT_RECORDS} records. "
            "Narrow the export with search or filters and try again."
        )
    return list(
        db.scalars(
            select(Submission)
            .where(total_where)
            .options(selectinload(Submission.submitted_by_user))
            .order_by(*order_by)
        )
    )


def render_csv(form: Form, submissions: list[Submission]) -> bytes:
    """Render records as UTF-8 CSV (BOM included)."""
    headers, rows = _rows(form, submissions)
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(headers)
    writer.writerows(rows)
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


_MULTI_SPACE = re.compile(r"\s+")


def _column_width(label: str) -> int:
    text = _MULTI_SPACE.sub(" ", label.strip())
    return min(max(len(text), 12), 40)


def render_xlsx(form: Form, submissions: list[Submission]) -> bytes:
    """Render records as a single ``Records`` worksheet (.xlsx)."""
    headers, _rows_display = _rows(form, submissions)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Records"
    worksheet.freeze_panes = "A2"

    for index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        cell.alignment = Alignment(vertical="center")
        worksheet.column_dimensions[cell.column_letter].width = _column_width(header)

    fields = _ordered_fields(form)
    first_data_column = len(EXPORT_META_COLUMNS) + 1

    for row_index, submission in enumerate(submissions, start=2):
        worksheet.cell(row=row_index, column=1, value=submission.id)
        submitter = submission.submitted_by_user
        submitter_name = (
            submitter.full_name
            if submitter and submitter.full_name
            else (submitter.username if submitter else "")
        )
        worksheet.cell(row=row_index, column=2, value=submitter_name or None)
        worksheet.cell(
            row=row_index,
            column=3,
            value=submission.submitted_at.replace(tzinfo=None) if submission.submitted_at else None,
        )
        for offset, field in enumerate(fields):
            column = first_data_column + offset
            worksheet.cell(
                row=row_index,
                column=column,
                value=_excel_value(field, submission.data.get(field.field_key)),
            )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _ordered_fields(form: Form):
    return sorted(form.fields, key=lambda f: f.sort_order)


def render_pdf(form: Form, submissions: list[Submission]) -> bytes:
    """Render records as a landscape PDF table with repeated headers.

    Long text is wrapped per cell (CJK, so any character can break). Single
    cells longer than :data:`_PDF_RECORD_CELL_LIMIT` characters are truncated
    with ``...``: an unbroken 10k-character blob in one narrow cell would
    otherwise grow a table row taller than the page and reportlab would abort
    the whole export with a LayoutError. CSV / XLSX / SQL always carry the
    complete value; the cap exists only so the PDF renderer is always safe.
    """
    headers, rows = _rows(form, submissions)

    buffer: BinaryIO = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        rightMargin=0.4 * inch,
        leftMargin=0.4 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
        title=f"Records — {form.name}",
        author="Dynamic Data Entry Platform",
    )
    available_width = landscape(letter)[0] - 0.8 * inch

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ExportTitle",
        parent=styles["Title"],
        fontSize=16,
        leading=20,
        alignment=TA_CENTER,
    )
    meta_style = ParagraphStyle(
        "ExportMeta",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        alignment=TA_CENTER,
        textColor=colors.grey,
    )
    cell_style = ParagraphStyle(
        "ExportCell",
        parent=styles["Normal"],
        fontSize=7,
        leading=9,
        wordWrap="CJK",
    )
    header_style = ParagraphStyle(
        "ExportHeaderCell",
        parent=styles["Normal"],
        fontSize=7.5,
        leading=9,
        textColor=colors.white,
        fontName="Helvetica-Bold",
    )

    story = [Paragraph(f"Records — {xml_escape(form.name)}", title_style)]
    exported_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    story.append(
        Paragraph(
            xml_escape(
                f"Form ID {form.id} · exported {exported_at} · "
                f"{len(submissions)} record(s)"
            ),
            meta_style,
        )
    )
    story.append(Spacer(1, 0.15 * inch))

    column_count = len(headers)
    column_width = available_width / max(column_count, 1)

    table_data: list[list[Paragraph]] = []
    table_data.append([Paragraph(xml_escape(header), header_style) for header in headers])
    for row in rows:
        table_data.append([Paragraph(xml_escape(_pdf_cell_text(cell)), cell_style) for cell in row])

    table = Table(
        table_data,
        colWidths=[column_width] * column_count,
        repeatRows=1,
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#334155")),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    return buffer.getvalue()


_PDF_RECORD_CELL_LIMIT = 400


def _pdf_cell_text(value: str, limit: int = _PDF_RECORD_CELL_LIMIT) -> str:
    """Truncate extremely long cells so a table row can always paginate.

    PDF cells wrap (CJK) normally; only a pathological single value would ever
    exceed the cap. CSV / XLSX / SQL carry the complete value -- this exists
    purely so the PDF renderer can never abort with a reportlab LayoutError.
    """
    if len(value) > limit:
        return value[: limit - 3].rstrip() + "..."
    return value


# --------------------------------------------------------------------------- #
# Phase 8: form-style PDF (individual + bulk)
#
# One portrait page per submission. Field labels and values come from the live
# form definition and are formatted with the same ``_display_value`` logic as
# every other export. There is no executable template system: the layout is
# generated entirely from the form definition and safe XML-escaped text.
#
# Non-Latin script runs (Devanagari / CJK) are wrapped in registered Unicode
# fonts when a matching font exists on the host; otherwise they fall back to
# the base font (reportlab never raises, though unsupported glyphs will not
# render accurately). This is the documented Unicode support ceiling.
# --------------------------------------------------------------------------- #

#: Script-range regexes used to detect Devanagari and CJK runs inside values.
_FORM_SCRIPT_SPLIT = re.compile(
    "(?P<devanagari>[\u0900-\u097f\ua8e0-\ua8ff\u1cd0-\u1cff]+)"
    "|(?P<cjk>[\u3000-\u303f\u3040-\u30ff\u3400-\u4dbf\u4e00-\u9fff"
    "\uf900-\ufaff\uac00-\ud7af]+)"
)

#: Directory roots where unicode TTF/TTC fonts are searched for. The whole
#: list is probed with ``os.path.isfile`` so machines without a font simply
#: degrade to the base Helvetica font instead of crashing.
_FORM_FONT_DIRS = tuple(
    dict.fromkeys(
        item
        for item in (
            os.path.join(os.environ.get("WINDIR") or "", "Fonts"),
            "C:/Windows/Fonts",
            "/usr/share/fonts/opentype/noto",
            "/usr/share/fonts/truetype/dejavu",
        )
        if item and not item.startswith("\\")
    )
)

#: Candidate font files per script. Lookup order matters; the first existing
#: file is registered (subfont 0) under the registered name.
_FORM_FONT_CANDIDATES = {
    "devanagari": ("DDEPDevanagari", ("nirmala.ttc", "NotoSansDevanagari-Regular.ttf", "DejaVuSans.ttf")),
    "cjk": ("DDEPCJK", ("msyh.ttc", "msyh.ttf", "simsun.ttc", "NotoSansCJK-Regular.ttc", "PingFang.ttc", "DejaVuSans.ttf")),
}

_UNICODE_FONT_CACHE: dict[str, str | None] = {}


def _unicode_font(script: str) -> str | None:
    """Return the registered font name for a script, or None if unavailable."""
    if script not in _UNICODE_FONT_CACHE:
        registered_name, candidates = _FORM_FONT_CANDIDATES[script]
        _UNICODE_FONT_CACHE[script] = None
        for directory in _FORM_FONT_DIRS:
            for filename in candidates:
                path = os.path.join(directory, filename)
                if not os.path.isfile(path):
                    continue
                try:
                    pdfmetrics.registerFont(TTFont(registered_name, path, subfontIndex=0))
                    _UNICODE_FONT_CACHE[script] = registered_name
                    break
                except Exception:
                    continue
            if _UNICODE_FONT_CACHE[script]:
                break
    return _UNICODE_FONT_CACHE.get(script)


def _font_marked_text(text: str) -> str:
    """XML-escape ``text`` and wrap non-Latin script runs in Unicode fonts.

    ASCII / Latin-1 runs stay in the base Helvetica font (so they remain
    directly extractable and consistent with the rest of the PDF output);
    Devanagari and CJK runs are wrapped in ``<font name="...">`` spans when a
    matching font was registered. All content is XML-escaped first, so the
    only markup emitted is the intentional ``<font>`` wrapper.
    """
    devanagari_font = _unicode_font("devanagari")
    cjk_font = _unicode_font("cjk")
    escaped = xml_escape(str(text))

    def wrap(font_name: str, chunk: str) -> str:
        return f'<font name="{font_name}">{chunk}</font>'

    pieces: list[str] = []
    pos = 0
    for match in _FORM_SCRIPT_SPLIT.finditer(escaped):
        pieces.append(escaped[pos : match.start()])
        devanagari, cjk = match.group("devanagari"), match.group("cjk")
        if devanagari and devanagari_font:
            pieces.append(wrap(devanagari_font, devanagari))
        elif cjk and cjk_font:
            pieces.append(wrap(cjk_font, cjk))
        else:
            pieces.append(match.group(0))
        pos = match.end()
    pieces.append(escaped[pos:])
    return "".join(pieces)


_FORM_LABEL_WIDTH_RATIO = 0.34
_FORM_MARGIN = 0.6 * inch


def _form_page_styles() -> tuple[ParagraphStyle, ParagraphStyle, ParagraphStyle, ParagraphStyle]:
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "FormRecordTitle",
        parent=styles["Title"],
        fontSize=17,
        leading=21,
        spaceAfter=2,
    )
    meta_style = ParagraphStyle(
        "FormRecordMeta",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        textColor=colors.grey,
        spaceAfter=8,
    )
    label_style = ParagraphStyle(
        "FormRecordFieldLabel",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        fontName="Helvetica-Bold",
        textColor=colors.HexColor("#334155"),
        wordWrap="CJK",
    )
    value_style = ParagraphStyle(
        "FormRecordFieldValue",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        wordWrap="CJK",
    )
    return title_style, meta_style, label_style, value_style


def _form_page_story(form: Form, submission: Submission) -> list:
    """Flowables for ONE submission's form-style page.

    Shared by the individual and bulk renderers so the layout cannot diverge.
    A title, a meta line, and a label/value table driven by the live form
    definition (ordered by ``sort_order``). Long and empty values are handled
    by ``_display_value`` + ``_pdf_cell_text`` in exactly the same way as the
    table PDF.
    """
    title_style, meta_style, label_style, value_style = _form_page_styles()

    submitter = submission.submitted_by_user
    submitter_name = (
        submitter.full_name
        if submitter and submitter.full_name
        else (submitter.username if submitter else "")
    )
    submitted_at = (
        submission.submitted_at.isoformat() if submission.submitted_at else ""
    )
    meta_text = (
        f"Form: {xml_escape(form.name)} (#{form.id}) \u00b7 "
        f"Submitted by {xml_escape(submitter_name or '')} \u00b7 {xml_escape(submitted_at)}"
    )

    available_width = letter[0] - 2 * _FORM_MARGIN
    label_width = available_width * _FORM_LABEL_WIDTH_RATIO
    value_width = available_width - label_width

    header_cell_style = ParagraphStyle(
        "FormRecordHeaderCell",
        parent=getSampleStyleSheet()["Normal"],
        fontSize=9.5,
        leading=12,
        fontName="Helvetica-Bold",
        textColor=colors.white,
    )

    field_rows: list[list[Paragraph]] = [
        [
            Paragraph("Field", header_cell_style),
            Paragraph("Value", header_cell_style),
        ]
    ]
    for field in _ordered_fields(form):
        stored = submission.data.get(field.field_key) if submission.data else None
        display = _pdf_cell_text(_display_value(field, stored))
        field_rows.append(
            [
                Paragraph(_font_marked_text(field.label), label_style),
                Paragraph(_font_marked_text(display), value_style),
            ]
        )

    table = Table(field_rows, colWidths=[label_width, value_width], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#334155")),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ]
        )
    )

    return [
        Paragraph(f"Record #{submission.id} \u2014 {_font_marked_text(form.name)}", title_style),
        Paragraph(meta_text, meta_style),
        Spacer(1, 0.1 * inch),
        table,
    ]


def _form_pdf_footer(canvas, doc) -> None:
    """Page footer for form-style PDFs draws only safe static text."""
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.grey)
    canvas.drawRightString(
        letter[0] - doc.rightMargin, 0.4 * inch, f"Page {doc.page} \u00b7 Dynamic Data Entry Platform"
    )
    canvas.restoreState()


def render_form_pdf(form: Form, submission: Submission) -> bytes:
    """Render ONE record as a portrait form-style PDF page."""
    buffer: BinaryIO = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=_FORM_MARGIN,
        leftMargin=_FORM_MARGIN,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
        title=f"Record #{submission.id} \u2014 {form.name}",
        author="Dynamic Data Entry Platform",
    )
    doc.build(
        _form_page_story(form, submission),
        onFirstPage=_form_pdf_footer,
        onLaterPages=_form_pdf_footer,
    )
    return buffer.getvalue()


def render_bulk_form_pdf(form: Form, submissions: list[Submission]) -> bytes:
    """Render MANY records as form-style PDF, exactly one submission per page.

    ``PageBreak`` is inserted between records so data from different records
    can never share a page; an empty result set yields a single blank page.
    """
    buffer: BinaryIO = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=_FORM_MARGIN,
        leftMargin=_FORM_MARGIN,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
        title=f"Records \u2014 {form.name}",
        author="Dynamic Data Entry Platform",
    )
    story: list = []
    for index, submission in enumerate(submissions):
        story.extend(_form_page_story(form, submission))
        if index < len(submissions) - 1:
            story.append(PageBreak())
    if not story:
        story.append(Paragraph("", _form_page_styles()[3]))
    doc.build(
        story,
        onFirstPage=_form_pdf_footer,
        onLaterPages=_form_pdf_footer,
    )
    return buffer.getvalue()


def export_filename(form: Form, format: str) -> str:
    """Safe attachment filename for the bulk export endpoint.

    The table PDF, form-style PDF and spreadsheet/SQL exports get distinct
    stems so two different artifacts never collide in a download folder.
    """
    slug = slugify_filename(form.name, form.id)
    if format == "pdf-form":
        return f"{slug}-records-form.pdf"
    if format == "pdf":
        return f"{slug}-records.pdf"
    return f"{slug}-records.{format}"


def individual_pdf_filename(form: Form, submission_id: int) -> str:
    """Safe attachment filename for an individual record's form PDF."""
    return f"{slugify_filename(form.name, form.id)}-record-{submission_id}.pdf"


def _sql_quote(value: str) -> str:
    """Single-quote a PostgreSQL string literal (safe against ``'``)."""
    return "'" + value.replace("'", "''") + "'"


def render_sql(form: Form, submissions: list[Submission]) -> bytes:
    """Render records as portable PostgreSQL INSERT statements.

    The output targets the real ``submissions`` table schema. It represents
    submission data for inspection / migration and is deliberately *not* a
    complete database backup (users, forms, fields, configuration, and other
    tables are not included).
    """
    exported_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    lines = [
        "-- Dynamic Data Entry Platform record export",
        f"-- Form: {form.name} (#{form.id})",
        f"-- Exported: {exported_at}",
        f"-- Records: {len(submissions)}",
        "-- Purpose: inspection / migration of submission data. This is NOT a",
        "-- complete database backup: other tables (users, forms, ...) are omitted.",
        "-- No user credentials, password hashes, or secrets are exported.",
        "SET standard_conforming_strings = on;",
        "BEGIN;",
        "INSERT INTO submissions (form_id, submitted_by, data, submitted_at)",
        "VALUES",
    ]
    parts: list[str] = []
    for index, submission in enumerate(submissions):
        data = json.dumps(submission.data, ensure_ascii=False, separators=(",", ":"))
        submitted_at = (
            submission.submitted_at.isoformat()
            if submission.submitted_at
            else datetime.now().isoformat()
        )
        values = (
            f"  ({int(submission.form_id)}, {int(submission.submitted_by)}, "
            f"{_sql_quote(data)}::jsonb, "
            f"{_sql_quote(submitted_at)})"
        )
        separator = ";" if index == len(submissions) - 1 else ","
        parts.append(values + separator)
    lines.extend(parts)
    lines.append("COMMIT;")
    return "\n".join(lines).encode("utf-8")


def generate_export(
    db: Session,
    form: Form,
    *,
    format: str,
    search: str | None,
    filters: str | None,
    sort_by: str | None,
    sort_order: str | None,
) -> tuple[bytes, str, str]:
    """Generate an export and return ``(content, media_type, extension)``.

    ``format`` is validated here (400-style error for anything outside
    ``EXPORT_FORMATS``). Search / filter / sort parameters are delegated to
    the Phase 5 query engine and may raise :class:`RecordQueryError`.
    """
    if format not in EXPORT_FORMATS:
        raise ExportError(
            f"Unsupported export format: {format}. "
            f"Supported formats: {', '.join(EXPORT_FORMATS)}"
        )

    submissions = _fetch_records(db, form, search, filters, sort_by, sort_order)

    if format == "csv":
        return render_csv(form, submissions), _MEDIA_TYPES["csv"], "csv"
    if format == "xlsx":
        return render_xlsx(form, submissions), _MEDIA_TYPES["xlsx"], "xlsx"
    if format == "pdf":
        return render_pdf(form, submissions), _MEDIA_TYPES["pdf"], "pdf"
    if format == "pdf-form":
        return render_bulk_form_pdf(form, submissions), _MEDIA_TYPES["pdf-form"], "pdf"
    return render_sql(form, submissions), _MEDIA_TYPES["sql"], "sql"